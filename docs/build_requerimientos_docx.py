# -*- coding: utf-8 -*-
"""Genera docs/Requerimientos_PLURALL.docx a partir de los datos de requerimientos.

Construye fichas de Análisis de Requerimiento (tabla Aspecto/Descripción) igual
al formato de image.png, para los RF y RNF del proyecto PLURALL.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- datos
ASPECTOS_RF = ["Objetivo del requisito", "Tipo de requisito", "Componentes involucrados",
               "Viabilidad técnica", "Riesgos potenciales", "Impacto en el sistema",
               "Dependencias", "Priorización"]
ASPECTOS_RNF = ["Objetivo del requisito", "Tipo de requisito", "Componentes involucrados",
                "Criterios de aceptación", "Riesgos potenciales", "Impacto en el sistema",
                "Dependencias", "Priorización"]

RF = [
 ("RF-01", "Autenticación Segura", [
  "Permitir que un usuario se registre, inicie sesión y recupere su contraseña de forma segura, usando el celular colombiano de 10 dígitos como identificador único de acceso.",
  "Funcional – Autenticación y gestión de acceso.",
  "Frontend (pantallas login / registro / recuperar), Backend (Supabase Auth, API de OTP y bloqueo), Base de datos (perfil_usuario, otp_recuperacion).",
  "Alta. Supabase Auth provee el manejo de sesiones y credenciales; el OTP y el bloqueo se implementan con tablas y API Routes propias.",
  "1. Ataques de fuerza bruta sobre cuentas.\n2. Correos electrónicos inválidos o duplicados en el registro.\n3. Fuga de sesión si el middleware no valida correctamente.",
  "Alto. Es la puerta de entrada a los cuatro portales; una falla compromete toda la plataforma.",
  "Supabase Auth, servicio de email (Resend) para OTP, middleware de sesión, catálogo de roles.",
  "Alta. Sin autenticación no es posible acceder a ninguna funcionalidad del sistema."],
  "Reglas de negocio: RN-01 (celular único), RN-02 (bloqueo tras 3 intentos por 30 min), validación de correo con dominio .com (SCRUM-33). | Jira: SCRUM-6, SCRUM-33 (Finalizado)."),

 ("RF-02", "Gestión de Perfil del Cliente", [
  "Permitir al cliente visualizar y editar sus datos de contacto y cambiar su contraseña, mostrando información real proveniente de la base de datos.",
  "Funcional – Gestión de datos personales.",
  "Frontend (pantalla de perfil), Backend (consulta/actualización), Base de datos (cliente, detalle_cliente, perfil_usuario).",
  "Alta. Operaciones CRUD estándar sobre tablas ya modeladas con RLS.",
  "1. Edición no autorizada de datos de otro usuario si RLS falla.\n2. Inconsistencia entre datos de contacto y datos de contrato.",
  "Medio. Afecta la experiencia del cliente y la exactitud de los datos de contacto para cobranza y notificaciones.",
  "Módulo de autenticación (RF-01), políticas RLS, base de datos de clientes.",
  "Media. Necesario para la operación pero no bloquea el flujo crediticio principal."],
  "Jira: SCRUM-11 (Finalizado)."),

 ("RF-03", "Solicitud de Crédito, Cuestionario y KYC", [
  "Guiar al cliente a través de un flujo multi-paso (cotizador → cuestionario de perfil → carga de documentos KYC → resumen) para radicar una solicitud de crédito.",
  "Funcional – Originación de crédito.",
  "Frontend (stepper de 4 pasos, formularios validados), Backend (POST /api/creditos/solicitar), Storage (bucket documentos_kyc), Base de datos (solicitud_credito, respuesta_cliente, documento_cliente).",
  "Alta. El flujo se apoya en componentes de formulario, validación de campos numéricos y carga de archivos a Supabase Storage.",
  "1. Envío de solicitudes incompletas si la validación por paso falla.\n2. Documentos KYC de baja calidad o no legibles.\n3. Solicitudes duplicadas si no se respeta el cooldown.",
  "Alto. Es el punto de partida de todo el ciclo de vida del crédito.",
  "Motor de scoring (RF-05), Storage, módulo de perfil (RF-02), calculadora de amortización.",
  "Alta. Núcleo funcional del negocio."],
  "Reglas de negocio: RN-03 (cooldown de 90 días), RN-04 (KYC obligatorio: cédula frontal + trasera). | Jira: SCRUM-16, SCRUM-34, SCRUM-35 (Finalizado)."),

 ("RF-04", "Gestión de Créditos y Documentos Legales", [
  "Mostrar el historial de créditos del usuario, con detalles como saldo, cuotas, estado de pago y documentos como contrato y paz y salvo.",
  "Funcional – Gestión de créditos y documentos.",
  "Frontend (visualización de historial de créditos), Backend (consulta de créditos y pagos, generación de PDF con Puppeteer), Base de datos (información financiera).",
  "Alta. La información de los créditos está organizada en la base de datos y es accesible para su consulta.",
  "1. Inexactitudes en la visualización del saldo o pagos realizados.\n2. Problemas con la descarga de documentos si no se implementa un sistema de almacenamiento adecuado.",
  "Alto. La correcta visualización y acceso a documentos legales es esencial para la transparencia.",
  "Módulo de créditos, servicio de documentos (paz y salvo, contratos, extracto de cuenta), base de datos financiera.",
  "Alta. Es fundamental que el usuario pueda acceder a su historial de créditos y descargar los documentos relacionados."],
  "Documentos generados: contrato, plan de pagos, extracto/estado de cuenta y paz y salvo (PDF vía Puppeteer). | Jira: SCRUM-20, SCRUM-24 (Finalizado)."),

 ("RF-05", "Motor de Scoring Crediticio Automático", [
  "Calcular automáticamente un puntaje crediticio a partir de las respuestas del cuestionario y factores ponderados, definiendo si la solicitud se aprueba, se rechaza o pasa a revisión manual.",
  "Funcional – Evaluación de riesgo.",
  "Backend (cálculo de score dentro de POST /api/creditos/solicitar), Base de datos (score, score_factor, cuestionario de preguntas).",
  "Alta. El modelo es determinista (respuestas × factores) y se ejecuta en el mismo endpoint de la solicitud.",
  "1. Umbrales mal calibrados que aprueben clientes de alto riesgo.\n2. Falta de trazabilidad del puntaje si no se persiste el detalle.",
  "Alto. Determina la decisión automática y la carga de trabajo de los analistas.",
  "Solicitud de crédito (RF-03), catálogo de preguntas y factores.",
  "Alta. Es el mecanismo que automatiza la mayor parte de las decisiones."],
  "Reglas de negocio: RN-09 (scoring visible para analista/supervisor). | Jira: SCRUM-18 (Finalizado)."),

 ("RF-06", "Decisión de Analista (Aprobar / Rechazar)", [
  "Permitir que un analista/supervisor revise las solicitudes en estudio y las apruebe o rechace, seleccionando un motivo del catálogo, con registro trazable de la decisión.",
  "Funcional – Gestión de decisiones.",
  "Frontend (panel de decisión con detalle de solicitud), Backend (PATCH /api/admin/solicitud/[id]), Base de datos (decision_solicitud, motivo_rechazo).",
  "Alta. Cambios de estado controlados sobre la solicitud con catálogo de motivos configurable.",
  "1. Cambio de una decisión ya tomada (inmutabilidad).\n2. Rechazos sin motivo registrado que afecten la auditoría.",
  "Alto. Habilita el paso de solicitud a desembolso.",
  "Scoring (RF-05), catálogo de motivos de rechazo, auditoría (RF-13).",
  "Alta. Cierra el flujo de originación previo al desembolso."],
  "Reglas de negocio: RN-05 (analista asignado), RN-06 (decisión inmutable). | Jira: SCRUM-19, SCRUM-23 (Finalizado)."),

 ("RF-07", "Pagos, Plan de Cuotas e Integración Wompi", [
  "Permitir al cliente consultar su plan de pagos y pagar cuotas a través del widget de Wompi, confirmando el pago mediante un webhook seguro e idempotente.",
  "Funcional – Procesamiento de pagos.",
  "Frontend (plan de pagos, widget Wompi), Backend (POST /api/pagos/iniciar, POST /api/webhooks/wompi), Rate limiting (Upstash Redis), Base de datos (plan_pago, cuota, pago).",
  "Media-Alta. Requiere firma HMAC SHA-256, validación de webhook e idempotencia para evitar dobles registros.",
  "1. Pagos duplicados si el webhook no es idempotente.\n2. Manipulación de montos si falla la validación de firma.\n3. Abuso del endpoint sin rate limiting.",
  "Alto. Afecta directamente el saldo del crédito y la salud de la cartera.",
  "Créditos y cuotas (RF-04), Wompi, Upstash Redis, notificaciones (RF-08).",
  "Alta. Es el mecanismo de recaudo de la plataforma."],
  "Reglas de negocio: RN-12 (cada pago genera registro con referencia Wompi), BD-03 (tasa congelada al desembolso). | Jira: SCRUM-22, SCRUM-25 (Finalizado)."),

 ("RF-08", "Notificaciones Transaccionales por Email", [
  "Enviar correos transaccionales al cliente ante eventos clave: aprobación, rechazo, confirmación de pago y recordatorio de vencimiento.",
  "Funcional – Comunicación con el usuario.",
  "Backend (/api/notificaciones/*, templates HTML), Servicio Resend, Base de datos (notificacion).",
  "Alta. Resend ofrece envío transaccional simple mediante API y plantillas HTML.",
  "1. Correos marcados como spam.\n2. Fallos silenciosos de envío sin reintento.\n3. Datos de contacto desactualizados.",
  "Medio. Mejora la experiencia y el recaudo, pero no bloquea las operaciones críticas.",
  "Resend, decisión de analista (RF-06), pagos (RF-07), perfil (RF-02).",
  "Media. Complementa los flujos principales."],
  "Jira: SCRUM-17 (Finalizado)."),

 ("RF-09", "Módulo de Cobranza para Asesores", [
  "Proveer al asesor herramientas para gestionar la cartera asignada: cola de cobranza, ficha de cliente en mora y registro de gestiones (llamadas, visitas, promesas de pago).",
  "Funcional – Gestión de cobranza.",
  "Frontend (portal /asesor), Backend (APIs de gestiones y asignación), Base de datos (gestion_cobranza, asignacion_asesor).",
  "Alta. Se apoya en las tablas de cartera y mora ya existentes.",
  "1. Asignaciones inconsistentes de clientes a asesores.\n2. Falta de trazabilidad de las gestiones registradas.",
  "Medio-Alto. Impacta la recuperación de cartera en mora.",
  "Detección de mora (RF-10), créditos y cuotas (RF-04), control de roles (RF-12).",
  "Media-Alta. Clave para la salud financiera del portafolio."],
  "Jira: SCRUM-29 (Finalizado)."),

 ("RF-10", "Detección Automática de Mora (Cron)", [
  "Ejecutar diariamente un proceso automático que marque cuotas vencidas, ponga créditos en mora y recupere créditos cuando ya no tengan cuotas vencidas.",
  "Funcional – Proceso automático (batch).",
  "Vercel Cron / pg_cron, Backend (GET /api/cron/mora autenticado con CRON_SECRET), Base de datos (cuota, credito).",
  "Alta. Vercel Cron dispara el endpoint a las 6 AM Colombia; alternativa con pg_cron disponible.",
  "1. Ejecución no autenticada del endpoint.\n2. Zona horaria mal configurada.\n3. Estados inconsistentes si el proceso se interrumpe.",
  "Alto. Mantiene actualizado el estado de mora que alimenta cobranza y analítica.",
  "Créditos y cuotas (RF-04), Vercel Cron, secreto de cron.",
  "Alta. Sostiene la exactitud del estado de la cartera."],
  "Reglas de negocio: RN-08 (cuota vencida automáticamente). | Jira: SCRUM-27 (Finalizado)."),

 ("RF-11", "Analítica y KPIs de Cartera", [
  "Presentar a supervisores y administradores indicadores de negocio (cartera, mora, funnel de originación, desembolsos) mediante dashboards con gráficas y drill-down.",
  "Funcional – Reportería y analítica.",
  "Frontend (dashboards con Recharts), Backend (/api/analitica, /api/reportes/export), Base de datos (vistas vw_kpi_cartera, vw_cohortes_desembolso, evento_funnel).",
  "Alta. Se apoya en vistas materializadas y componentes de gráficas.",
  "1. Métricas incorrectas por vistas desactualizadas.\n2. Consultas pesadas que degraden el rendimiento.",
  "Medio-Alto. Soporta la toma de decisiones gerenciales.",
  "Datos de créditos, pagos y funnel; vistas de reportería; control de roles (RF-12).",
  "Media-Alta. Aporta valor de gestión sin bloquear la operación transaccional."],
  "Jira: SCRUM-28, SCRUM-30 (Finalizado)."),

 ("RF-12", "Panel Administrativo y Control de Acceso por Roles", [
  "Ofrecer portales diferenciados (cliente, asesor, supervisor, admin) con acceso restringido según el rol del usuario, tanto en la interfaz como en los datos.",
  "Funcional – Autorización y administración.",
  "Frontend (layouts por rol, directorio de clientes), Middleware de sesión, Base de datos (rol, perfil_usuario, políticas RLS en todas las tablas).",
  "Alta. RLS de Supabase + verificación en layout garantizan el aislamiento por rol.",
  "1. Escalamiento de privilegios si una política RLS es incompleta.\n2. Exposición de datos entre clientes.",
  "Alto. Es la base de la seguridad de acceso a la información.",
  "Autenticación (RF-01), modelo de datos con RLS, middleware.",
  "Alta. Transversal a todos los módulos."],
  "Jira: SCRUM-13, SCRUM-21 (Finalizado)."),

 ("RF-13", "Auditoría y Trazabilidad del Sistema", [
  "Registrar quién realizó cada acción sensible, cuándo y desde qué IP, así como el historial de cambios de estado, para garantizar la trazabilidad.",
  "Funcional – Auditoría y control.",
  "Backend (registro de eventos), Base de datos (auditoria, historial_estado, config_sistema).",
  "Alta. Registro de eventos mediante inserciones controladas y/o triggers.",
  "1. Registros incompletos que impidan reconstruir un incidente.\n2. Crecimiento no controlado de la tabla de auditoría.",
  "Medio-Alto. Esencial para control interno y cumplimiento.",
  "Decisiones (RF-06), pagos (RF-07), autenticación (RF-01).",
  "Media-Alta. Requerido para operación confiable y auditable."],
  "Jira: SCRUM-12 (Finalizado)."),

 ("RF-14", "PQR y Asistencia al Usuario", [
  "Permitir al cliente radicar Peticiones, Quejas o Reclamos (PQR) y ofrecer una sesión de ayuda o chatbot de soporte dentro de la plataforma.",
  "Funcional – Atención al usuario.",
  "Frontend (formulario PQR, componente de ayuda/chat), Backend (/api/pqr), Base de datos (pqr).",
  "Alta para PQR (CRUD estándar); media para el chatbot según el alcance de la asistencia.",
  "1. PQR sin seguimiento ni SLA.\n2. Expectativas del chatbot superiores a su alcance real.",
  "Medio. Mejora la experiencia y el cumplimiento de atención al cliente.",
  "Autenticación (RF-01), notificaciones (RF-08).",
  "Media. Complementa la experiencia del cliente."],
  "Jira: SCRUM-31 (Finalizado)."),
]

RNF = [
 ("RNF-01", "Seguridad", [
  "Garantizar la confidencialidad, integridad y disponibilidad de los datos mediante autenticación robusta, autorización por RLS y validación de integridad en pagos.",
  "No funcional – Seguridad.",
  "Supabase Auth y RLS, middleware, HMAC SHA-256 (Wompi), headers HTTP (HSTS, X-Frame-Options, X-Content-Type-Options), rate limiting.",
  "100% de tablas con políticas RLS; webhooks validados por firma; bloqueo tras 3 intentos fallidos; headers de seguridad en todas las respuestas.",
  "Fuerza bruta, escalamiento de privilegios, manipulación de webhooks, secuestro de sesión.",
  "Crítico. Afecta a toda la plataforma y a la confianza del usuario.",
  "RF-01, RF-07, RF-12.",
  "Alta."], ""),
 ("RNF-02", "Rendimiento", [
  "Mantener tiempos de respuesta aceptables en la navegación y las consultas, incluso en dashboards con agregaciones.",
  "No funcional – Rendimiento.",
  "Server Components (SSR) de Next.js, vistas materializadas, índices en base de datos, Edge/CDN de Vercel.",
  "Páginas principales cargan en menos de ~3 s en condiciones normales; los KPIs se sirven desde vistas para evitar consultas costosas en tiempo real.",
  "Consultas no indexadas, vistas desactualizadas, sobrecarga en horas pico.",
  "Medio-Alto. Afecta la experiencia de todos los roles.",
  "RF-11, modelo de datos.",
  "Media-Alta."], ""),
 ("RNF-03", "Disponibilidad y Despliegue", [
  "Asegurar que la plataforma esté disponible de forma continua mediante hosting serverless y despliegue automatizado.",
  "No funcional – Disponibilidad.",
  "Vercel (Serverless + Edge + Cron), Supabase gestionado, integración continua desde el repositorio.",
  "Despliegues reproducibles desde Git; cron de mora ejecutándose diariamente; variables de entorno gestionadas por entorno.",
  "Caídas de servicios externos (Wompi, Resend), errores de configuración de entorno.",
  "Alto. La indisponibilidad detiene la operación.",
  "RF-10, infraestructura Vercel/Supabase.",
  "Alta."], ""),
 ("RNF-04", "Usabilidad y Diseño Responsivo", [
  "Ofrecer una interfaz clara, consistente y utilizable desde dispositivos móviles y de escritorio, con un sistema de diseño unificado.",
  "No funcional – Usabilidad.",
  "Tailwind CSS, componentes reutilizables, sidebar responsivo, sistema de diseño global.",
  "Navegación funcional en móvil y escritorio; componentes visibles y consistentes; mejoras de visibilidad aplicadas (SCRUM-32).",
  "Inconsistencias visuales entre portales, componentes poco visibles en algunas resoluciones.",
  "Medio. Afecta adopción y satisfacción del usuario.",
  "Todos los módulos con interfaz.",
  "Media."], "Jira: SCRUM-32 (Finalizado)."),
 ("RNF-05", "Escalabilidad", [
  "Permitir el crecimiento del número de usuarios, créditos y transacciones sin rediseñar la arquitectura.",
  "No funcional – Escalabilidad.",
  "Arquitectura serverless (Vercel), PostgreSQL gestionado (Supabase), rate limiting distribuido (Upstash Redis).",
  "La aplicación escala horizontalmente en funciones serverless; la base de datos soporta el crecimiento con índices y vistas.",
  "Límites de plan gratuito, cuellos de botella en la base de datos.",
  "Medio. Relevante a mediano plazo.",
  "Infraestructura, RNF-02.",
  "Media."], ""),
 ("RNF-06", "Mantenibilidad", [
  "Facilitar la evolución del sistema mediante código organizado, migraciones versionadas y documentación por módulo y sprint.",
  "No funcional – Mantenibilidad.",
  "Estructura de repositorio (apps/web, db, docs), migraciones SQL numeradas, componentes reutilizables.",
  "Migraciones aplicables en orden; documentación de sprints y módulos actualizada; convenciones de código consistentes.",
  "Deuda técnica, migraciones fuera de orden, base de datos sin normalizar (SCRUM-36).",
  "Medio. Afecta la velocidad de desarrollo futuro.",
  "Todo el proyecto.",
  "Media."], "Jira: SCRUM-36 (Por hacer)."),
 ("RNF-07", "Monitoreo y Observabilidad", [
  "Detectar y diagnosticar errores en producción mediante captura de excepciones y repetición de sesiones.",
  "No funcional – Observabilidad.",
  "Sentry (errores + Session Replay), tabla de auditoría, logs de Vercel.",
  "Errores capturados en Sentry omitiendo 401/403 esperados; posibilidad de reconstruir incidentes.",
  "Ruido de errores esperados, falta de alertas oportunas.",
  "Medio-Alto. Reduce el tiempo de detección y resolución de fallos.",
  "RF-13, infraestructura.",
  "Media-Alta."], ""),
 ("RNF-08", "Cumplimiento Legal (Habeas Data)", [
  "Cumplir con la normativa colombiana de tratamiento de datos personales (Habeas Data), publicando la política de tratamiento y garantizando el consentimiento.",
  "No funcional – Legal / Cumplimiento.",
  "Documento legal público de tratamiento de datos (/legal), gestión de consentimiento, seguridad de datos personales.",
  "Política de tratamiento accesible públicamente (PDF); datos personales protegidos por RLS y cifrado en tránsito (HTTPS).",
  "Incumplimiento normativo, exposición de datos sensibles.",
  "Alto. Riesgo legal y reputacional.",
  "RNF-01, RF-01.",
  "Alta."], ""),
 ("RNF-09", "Compatibilidad", [
  "Garantizar el funcionamiento correcto en navegadores modernos y en distintos tamaños de pantalla.",
  "No funcional – Compatibilidad.",
  "Frontend Next.js, Tailwind CSS, estándares web.",
  "Funciona en versiones recientes de Chrome, Firefox, Edge y Safari, en móvil y escritorio.",
  "Comportamientos divergentes entre navegadores.",
  "Medio. Afecta el alcance de usuarios.",
  "RNF-04.",
  "Media."], ""),
 ("RNF-10", "Integridad y Consistencia de Datos", [
  "Mantener la coherencia de la información financiera y de estados mediante restricciones, triggers y procesos idempotentes.",
  "No funcional – Integridad de datos.",
  "Restricciones e índices en PostgreSQL, triggers, webhook idempotente de pagos, normalización de tablas.",
  "Un solo crédito ACTIVO por cliente; pagos sin duplicados; estados de crédito y cuota consistentes tras cada proceso.",
  "Registros duplicados, estados inconsistentes, datos sin normalizar (SCRUM-36 pendiente).",
  "Alto. La exactitud financiera depende de esta integridad.",
  "RF-07, RF-10, modelo de datos.",
  "Alta."], "Reglas de negocio: RN-07 (un solo crédito ACTIVO), RN-12, BD-03. | Jira: SCRUM-36 (Por hacer)."),
]

# ---------------------------------------------------------------- helpers
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
HEADBG = "1F3A5F"
ASPBG = "EAEFF5"

def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align: p.alignment = align
    for i, line in enumerate(str(text).split("\n")):
        if i > 0:
            r = p.add_run(); r.add_break()
        else:
            r = p.add_run(line); r.font.bold = bold; r.font.size = Pt(size)
            if color: r.font.color.rgb = color
            continue
        r = p.add_run(line); r.font.bold = bold; r.font.size = Pt(size)
        if color: r.font.color.rgb = color

def add_ficha(doc, codigo, nombre, valores, aspectos, nota):
    h = doc.add_paragraph()
    h.space_before = Pt(6)
    r = h.add_run(f"Análisis Requerimiento — {codigo} · {nombre}")
    r.font.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY

    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width = Inches(1.9); t.columns[1].width = Inches(4.6)
    hdr = t.rows[0].cells
    set_cell_text(hdr[0], "Aspecto", bold=True, color=RGBColor(0xFF,0xFF,0xFF), size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(hdr[1], "Descripción", bold=True, color=RGBColor(0xFF,0xFF,0xFF), size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_bg(hdr[0], HEADBG); set_cell_bg(hdr[1], HEADBG)
    for asp, val in zip(aspectos, valores):
        row = t.add_row().cells
        row[0].width = Inches(1.9); row[1].width = Inches(4.6)
        set_cell_text(row[0], asp, bold=True, size=10)
        set_cell_bg(row[0], ASPBG)
        set_cell_text(row[1], val, size=10)
    if nota:
        np = doc.add_paragraph()
        rr = np.add_run(nota); rr.font.italic = True; rr.font.size = Pt(8.5)
        rr.font.color.rgb = RGBColor(0x55,0x55,0x55)
    doc.add_paragraph()

def heading(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.bold = True
    r.font.size = Pt(16 if level == 1 else 13)
    r.font.color.rgb = NAVY

# ---------------------------------------------------------------- documento
doc = Document()
style = doc.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(10.5)

# Portada
for _ in range(4): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PLURALL"); r.font.bold = True; r.font.size = Pt(40); r.font.color.rgb = NAVY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Documento de Requerimientos del Sistema"); r.font.size = Pt(18); r.font.color.rgb = RGBColor(0x44,0x44,0x44)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Requerimientos Funcionales y No Funcionales"); r.font.size = Pt(13); r.font.italic = True
for _ in range(3): doc.add_paragraph()
for line in ["Plataforma Fintech de Gestión de Créditos",
             "Proyecto Formativo SENA — Tecnólogo en Análisis y Desarrollo de Software (ADSO)",
             "Instructor: Jefry González",
             "Equipo: Wilder Feo · Camila Avila · Juan Cetina · Juan Maurello · Jeferson Lugo",
             "Tablero de referencia: Jira — Proyecto SCRUM (plucardsito.atlassian.net)"]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(line); r.font.size = Pt(11)
doc.add_page_break()

# Introducción
heading(doc, "1. Introducción")
doc.add_paragraph(
 "Este documento consolida los requerimientos funcionales (RF) y no funcionales (RNF) de la "
 "plataforma fintech PLURALL, elaborado a partir de las historias, épicas y tareas del tablero "
 "de Jira del proyecto SCRUM. Cada requerimiento funcional se documenta bajo la ficha de Análisis "
 "de Requerimiento con los aspectos: objetivo, tipo, componentes involucrados, viabilidad técnica, "
 "riesgos potenciales, impacto en el sistema, dependencias y priorización.")

# Índice RF
heading(doc, "2. Requerimientos Funcionales", level=2)
idx = doc.add_table(rows=1, cols=3); idx.style = "Table Grid"
hc = idx.rows[0].cells
for i, t in enumerate(["Código", "Nombre", "Jira"]):
    set_cell_text(hc[i], t, bold=True, color=RGBColor(0xFF,0xFF,0xFF), size=10, align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_bg(hc[i], HEADBG)
rf_jira = {"RF-01":"SCRUM-6, 33","RF-02":"SCRUM-11","RF-03":"SCRUM-16, 34, 35","RF-04":"SCRUM-20, 24",
 "RF-05":"SCRUM-18","RF-06":"SCRUM-19, 23","RF-07":"SCRUM-22, 25","RF-08":"SCRUM-17","RF-09":"SCRUM-29",
 "RF-10":"SCRUM-27","RF-11":"SCRUM-28, 30","RF-12":"SCRUM-13, 21","RF-13":"SCRUM-12","RF-14":"SCRUM-31"}
for cod, nom, _, _ in RF:
    c = idx.add_row().cells
    set_cell_text(c[0], cod, bold=True, size=9); set_cell_text(c[1], nom, size=9); set_cell_text(c[2], rf_jira[cod], size=9)
doc.add_paragraph()

for cod, nom, vals, nota in RF:
    add_ficha(doc, cod, nom, vals, ASPECTOS_RF, nota)

doc.add_page_break()
heading(doc, "3. Requerimientos No Funcionales", level=2)
doc.add_paragraph()
for cod, nom, vals, nota in RNF:
    add_ficha(doc, cod, nom, vals, ASPECTOS_RNF, nota)

# Trazabilidad
doc.add_page_break()
heading(doc, "4. Trazabilidad Requerimiento ↔ Jira", level=2)
tr = doc.add_table(rows=1, cols=3); tr.style = "Table Grid"
hc = tr.rows[0].cells
for i, t in enumerate(["Requerimiento", "Tareas / Épicas Jira", "Estado"]):
    set_cell_text(hc[i], t, bold=True, color=RGBColor(0xFF,0xFF,0xFF), size=10, align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_bg(hc[i], HEADBG)
traz = [
 ("RF-01","SCRUM-6, SCRUM-33","Finalizado"),("RF-02","SCRUM-11","Finalizado"),
 ("RF-03","SCRUM-16, SCRUM-34, SCRUM-35","Finalizado"),("RF-04","SCRUM-20, SCRUM-24","Finalizado"),
 ("RF-05","SCRUM-18","Finalizado"),("RF-06","SCRUM-19, SCRUM-23","Finalizado"),
 ("RF-07","SCRUM-22, SCRUM-25","Finalizado"),("RF-08","SCRUM-17","Finalizado"),
 ("RF-09","SCRUM-29","Finalizado"),("RF-10","SCRUM-27","Finalizado"),
 ("RF-11","SCRUM-28, SCRUM-30","Finalizado"),("RF-12","SCRUM-13, SCRUM-21","Finalizado"),
 ("RF-13","SCRUM-12","Finalizado"),("RF-14","SCRUM-31","Finalizado"),
 ("RNF-04","SCRUM-32","Finalizado"),("RNF-06 / RNF-10","SCRUM-36 (normalizar BD)","Por hacer"),
]
for a,b,c in traz:
    row = tr.add_row().cells
    set_cell_text(row[0], a, bold=True, size=9); set_cell_text(row[1], b, size=9); set_cell_text(row[2], c, size=9)
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Épicas (Sprints): SCRUM-1 Fundación y Autenticación · SCRUM-2 Dashboard y Perfil · "
 "SCRUM-3 Solicitud de Crédito y KYC · SCRUM-4 Pagos, Cuotas y Decisiones · SCRUM-5 Producción, Auditoría y Finalización.")
r.font.size = Pt(9); r.font.italic = True

out = r"C:\Users\USUARIO\Documents\Estudio\Plurall\docs\Requerimientos_PLURALL.docx"
doc.save(out)
print("Guardado:", out)
